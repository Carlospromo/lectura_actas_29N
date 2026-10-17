"""Fuentes de contexto: mesas 2023 y censo de consulta 2026.

No crea mesas operativas, asignaciones, candidaturas ni resultados de 2026.
"""

import argparse
import hashlib
import re
import uuid
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

from app.catalog import PROVINCES, _territory
from app.db import Session, all_rows, one
from app.migrate import migrate

HISTORY_HEADERS = (
    "cod.mesa", "ANYO", "MES", "COD.CCAA", "CCAA", "COD.PROV", "PROVINCIA",
    "COD.MUN", "MUNICIPIO", "DISTRITO", "SECCION", "MESA", "CENSO",
)


def _code(value, width, label):
    if isinstance(value, bool) or value is None or not re.fullmatch(r"\d+", str(value)):
        raise ValueError(f"Código {label} inválido: {value!r}")
    number = int(value)
    if number < 0 or len(str(number)) > width:
        raise ValueError(f"Código {label} fuera de rango: {value!r}")
    return f"{number:0{width}d}"


def parse_history(rows):
    iterator = iter(rows)
    header = next(iterator)
    if tuple(header[:13]) != HISTORY_HEADERS:
        raise ValueError("Cabecera de MESAS distinta de la fuente esperada")
    records = []
    seen_raw, seen_identity = set(), set()
    communities, provinces, municipalities, exterior = {}, {}, {}, {}
    for line, row in enumerate(iterator, 2):
        if not row or all(v is None for v in row):
            continue
        raw = str(row[0]).strip()
        if (row[1], row[2]) != (2023, 7):
            raise ValueError(f"Fila {line}: periodo histórico inesperado")
        cc, p, m = _code(row[3], 2, "CCAA"), _code(row[5], 2, "provincia"), _code(row[7], 3, "municipio")
        district, section = _code(row[9], 2, "distrito"), _code(row[10], 3, "sección")
        label = str(row[11]).strip().upper()
        names = [row[i] for i in (4, 6, 8)]
        if p not in PROVINCES or not 1 <= int(cc) <= 19 or not label or not all(isinstance(v, str) and v.strip() for v in names):
            raise ValueError(f"Fila {line}: identidad territorial inválida")
        if raw != f"{cc}{p}{m}{district}{section}{label}":
            raise ValueError(f"Fila {line}: cod.mesa no coincide con sus componentes")
        if isinstance(row[12], bool) or not isinstance(row[12], int) or row[12] < 0:
            raise ValueError(f"Fila {line}: censo histórico inválido")
        identity = (p, m, district, section, label)
        if raw in seen_raw or identity in seen_identity:
            raise ValueError(f"Fila {line}: mesa histórica duplicada")
        seen_raw.add(raw); seen_identity.add(identity)
        if cc in communities and communities[cc] != names[0]:
            raise ValueError(f"Fila {line}: comunidad con nombre contradictorio")
        if p in provinces and provinces[p] != (cc, names[1]):
            raise ValueError(f"Fila {line}: provincia contradictoria")
        communities[cc] = names[0]
        provinces[p] = (cc, names[1])
        if m == "999":
            if district != "09" or section != "000" or label != "U":
                raise ValueError(f"Fila {line}: registro exterior con estructura inesperada")
            exterior[p] = names[2]
        else:
            key = p + m
            if key in municipalities and municipalities[key] != (p, names[2]):
                raise ValueError(f"Fila {line}: municipio contradictorio")
            municipalities[key] = (p, names[2])
        records.append({"raw_code": raw, "community_code": cc, "community_name": names[0], "province_code": p, "province_name": names[1], "municipality_code": m, "municipality_name": names[2], "district": district, "section": section, "table_label": label, "is_exterior": m == "999", "historical_census": row[12]})
    if set(communities) != {f"{n:02d}" for n in range(1, 20)} or set(provinces) != PROVINCES:
        raise ValueError("El fichero histórico no cubre las 19 comunidades/ciudades autónomas y 52 provincias")
    if set(exterior) != PROVINCES:
        raise ValueError("Faltan registros 999 de voto exterior por provincia")
    return records, communities, provinces, municipalities, exterior


def parse_census(rows):
    rows = list(rows)
    if len(rows) < 61 or "2026" not in str(rows[1][0]) or "Censo de consulta" not in str(rows[3][0]):
        raise ValueError("El archivo no identifica el censo de consulta de 2026")
    reported = rows[8]
    if str(reported[0]).strip() != "Total general":
        raise ValueError("Falta el total general")
    result = []
    seen = set()
    for line, row in enumerate(rows[9:], 10):
        if not row or all(v is None for v in row):
            continue
        p = _code(row[0], 2, "provincia")
        if p not in PROVINCES or p in seen or not isinstance(row[1], str) or not row[1].strip():
            raise ValueError(f"Fila {line}: provincia inválida o duplicada")
        spain, abroad, total = row[2:5]
        if any(isinstance(v, bool) or not isinstance(v, int) or v < 0 for v in (spain, abroad, total)) or total != spain + abroad:
            raise ValueError(f"Fila {line}: totales censales inconsistentes")
        seen.add(p)
        result.append({"province_code": p, "province_name": row[1].strip(), "residents_spain": spain, "residents_abroad": abroad, "total": total})
    if seen != PROVINCES or tuple(sum(r[k] for r in result) for k in ("residents_spain", "residents_abroad", "total")) != tuple(reported[2:5]):
        raise ValueError("El censo no concilia las 52 provincias con el total general")
    return result, {"residents_spain": reported[2], "residents_abroad": reported[3], "total": reported[4], "status": "consulta"}


def _hash(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _existing_or_insert(db, kind, version, path, year, label, count):
    sha = _hash(path)
    existing = one(db, "SELECT * FROM reference_imports WHERE kind=:kind AND source_sha256=:sha", {"kind": kind, "sha": sha})
    if existing:
        return existing["id"], False
    if one(db, "SELECT id FROM reference_imports WHERE kind=:kind AND version=:version", {"kind": kind, "version": version}):
        raise ValueError("La versión ya existe con otro contenido")
    import_id = str(uuid.uuid4())
    db.execute(text("INSERT INTO reference_imports(id,kind,version,source_name,source_sha256,source_year,source_label,row_count) VALUES (:id,:kind,:version,:name,:sha,:year,:label,:count)"), {"id": import_id, "kind": kind, "version": version, "name": path.name, "sha": sha, "year": year, "label": label, "count": count})
    return import_id, True


def import_history(path: Path, version: str):
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        records, communities, provinces, municipalities, exterior = parse_history(workbook["MESAS"].values)
    finally:
        workbook.close()
    with Session.begin() as db:
        import_id, new = _existing_or_insert(db, "historical_tables", version, path, 2023, "Mesas generales 2023-07; referencia histórica, sin locales", len(records))
        if not new:
            return {"id": import_id, "rows": len(records), "already_imported": True}
        cc_ids = {code: _territory(db, "community", code, name) for code, name in communities.items()}
        p_ids = {code: _territory(db, "province", code, name, cc_ids[cc]) for code, (cc, name) in provinces.items()}
        for code, (province, name) in municipalities.items():
            _territory(db, "municipality", code, name, p_ids[province])
        for province, name in exterior.items():
            _territory(db, "exterior", province + "999", name, p_ids[province])
        sql = text("INSERT INTO historical_table_references(import_id,raw_code,community_code,community_name,province_code,province_name,municipality_code,municipality_name,district,section,table_label,is_exterior,historical_census) VALUES (:import_id,:raw_code,:community_code,:community_name,:province_code,:province_name,:municipality_code,:municipality_name,:district,:section,:table_label,:is_exterior,:historical_census)")
        for start in range(0, len(records), 1000):
            db.execute(sql, [{"import_id": import_id, **r} for r in records[start:start + 1000]])
    return {"id": import_id, "rows": len(records), "communities": len(communities), "provinces": len(provinces), "municipalities": len(municipalities), "exterior_records": len(exterior), "already_imported": False}


def import_census(path: Path, version: str):
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        records, totals = parse_census(workbook["Electores por provincias"].values)
    finally:
        workbook.close()
    with Session.begin() as db:
        import_id, new = _existing_or_insert(db, "census_provinces", version, path, 2026, "Censo de consulta; referencia provisional", len(records))
        if not new:
            return {"id": import_id, "rows": len(records), "already_imported": True, **totals}
        db.execute(text("INSERT INTO census_province_references(import_id,province_code,province_name,residents_spain,residents_abroad,total) VALUES (:import_id,:province_code,:province_name,:residents_spain,:residents_abroad,:total)"), [{"import_id": import_id, **r} for r in records])
    return {"id": import_id, "rows": len(records), "already_imported": False, **totals}


def main():
    parser = argparse.ArgumentParser(description="Importa fuentes históricas y censales sin activar mesas de 2026")
    sub = parser.add_subparsers(dest="command", required=True)
    h = sub.add_parser("history")
    h.add_argument("xlsx", type=Path)
    h.add_argument("--version", required=True)
    c = sub.add_parser("census")
    c.add_argument("xlsx", type=Path)
    c.add_argument("--version", required=True)
    sub.add_parser("list")
    args = parser.parse_args()
    migrate()
    if args.command == "list":
        with Session() as db:
            print(all_rows(db, "SELECT id,kind,version,source_name,source_sha256,source_year,source_label,row_count,imported_at FROM reference_imports ORDER BY imported_at DESC"))
        return
    if args.command == "history":
        print(import_history(args.xlsx, args.version))
    else:
        print(import_census(args.xlsx, args.version))


if __name__ == "__main__":
    main()
